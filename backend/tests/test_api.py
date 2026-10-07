import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app import main
from app.providers import ProviderError

client = TestClient(main.app)

LETTER = "Dear Hiring Manager,\n\nI am excited — truly — to apply.\n\nSincerely,\nAda Lovelace"


class FakeProvider:
    name = "Fake"
    default_model = "fake-1"
    env_key = "FAKE_API_KEY"

    def __init__(self, error: ProviderError | None = None):
        self.error = error
        self.calls = []

    def generate(self, docs, instructions, api_key, model):
        self.calls.append((docs, instructions, api_key, model))
        if self.error:
            raise self.error
        return LETTER


@pytest.fixture
def fake(monkeypatch):
    provider = FakeProvider()
    monkeypatch.setitem(main.PROVIDERS, "fake", provider)
    monkeypatch.delenv("FAKE_API_KEY", raising=False)
    return provider


def _docx_bytes(*paragraphs: str) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_providers_lists_all():
    ids = {p["id"] for p in client.get("/api/providers").json()}
    assert {"anthropic", "openai", "gemini"} <= ids


def test_generate_with_docx_resume_and_text_jd(fake):
    r = client.post(
        "/api/generate",
        files={"resume": ("cv.docx", _docx_bytes("Ada Lovelace", "Analyst"), "application/octet-stream")},
        data={"jd_text": "Senior engineer", "provider": "fake", "api_key": "user-key"},
    )
    assert r.status_code == 200, r.text
    assert r.json() == {"cover_letter": LETTER, "provider": "fake", "model": "fake-1"}
    docs, _, key, model = fake.calls[0]
    assert [d.title for d in docs] == ["resume", "job_description"]
    assert "Ada Lovelace" in docs[0].text
    assert key == "user-key" and model == "fake-1"


def test_generate_uses_server_key_when_user_key_blank(fake, monkeypatch):
    monkeypatch.setenv("FAKE_API_KEY", "server-key")
    r = client.post(
        "/api/generate",
        files={"resume": ("cv.txt", b"Ada", "text/plain")},
        data={"jd_text": "Role", "provider": "fake", "model": "fake-2"},
    )
    assert r.status_code == 200, r.text
    assert fake.calls[0][2:] == ("server-key", "fake-2")


def test_generate_requires_key(fake):
    r = client.post(
        "/api/generate",
        files={"resume": ("cv.txt", b"Ada", "text/plain")},
        data={"jd_text": "Role", "provider": "fake"},
    )
    assert r.status_code == 400
    assert "API key" in r.json()["detail"]


def test_generate_requires_jd(fake):
    r = client.post(
        "/api/generate",
        files={"resume": ("cv.txt", b"Ada", "text/plain")},
        data={"provider": "fake", "api_key": "k"},
    )
    assert r.status_code == 400


def test_generate_rejects_unsupported_file(fake):
    r = client.post(
        "/api/generate",
        files={"resume": ("cv.png", b"\x89PNG", "image/png")},
        data={"jd_text": "Role", "provider": "fake", "api_key": "k"},
    )
    assert r.status_code == 400
    assert "unsupported" in r.json()["detail"]


def test_generate_unknown_provider():
    r = client.post(
        "/api/generate",
        files={"resume": ("cv.txt", b"Ada", "text/plain")},
        data={"jd_text": "Role", "provider": "nope", "api_key": "k"},
    )
    assert r.status_code == 400


def test_provider_error_is_propagated(monkeypatch):
    monkeypatch.setitem(main.PROVIDERS, "fake", FakeProvider(ProviderError(401, "bad key")))
    r = client.post(
        "/api/generate",
        files={"resume": ("cv.txt", b"Ada", "text/plain")},
        data={"jd_text": "Role", "provider": "fake", "api_key": "k"},
    )
    assert r.status_code == 401
    assert r.json()["detail"] == "bad key"


def test_export_pdf_contains_text():
    r = client.post("/api/export", json={"text": LETTER, "format": "pdf"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    assert "Ada Lovelace" in text and "Dear Hiring Manager" in text


def test_export_docx_keeps_paragraphs():
    r = client.post("/api/export", json={"text": LETTER, "format": "docx"})
    assert r.status_code == 200
    paras = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert paras[0] == "Dear Hiring Manager,"
    assert paras[-1] == "Sincerely,\nAda Lovelace"


def test_export_pdf_escapes_markup():
    r = client.post("/api/export", json={"text": "A <b> & C", "format": "pdf"})
    assert r.status_code == 200
