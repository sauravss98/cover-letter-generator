"""Render cover letter text to DOCX or PDF bytes. Paragraphs are separated by blank lines."""

import io
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text.strip()) if p.strip()]


def to_docx(text: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for para in _paragraphs(text):
        p = doc.add_paragraph()
        lines = para.split("\n")
        for i, line in enumerate(lines):
            run = p.add_run(line)
            if i < len(lines) - 1:
                run.add_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(text: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="Cover Letter",
    )
    style = ParagraphStyle(
        "body", fontName="Helvetica", fontSize=11, leading=15, spaceAfter=11
    )
    story = [
        Paragraph(escape(para).replace("\n", "<br/>"), style)
        for para in _paragraphs(text)
    ]
    doc.build(story)
    return buf.getvalue()
